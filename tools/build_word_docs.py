"""Convert the Markdown reports in reports/ to styled Word (.docx) files.

Uses pandoc (Markdown -> Word, LaTeX maths -> native Word equations), then
python-docx to style tables (borders, shaded header row) and fonts.

Run from the project root:
    python tools/build_word_docs.py
"""

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPORTS = Path("reports")
DOCS = {
    # markdown file: (output, body font, heading colour, add table of contents?)
    "methodology_journal.md": ("methodology_journal.docx", "Cambria", RGBColor(0x1F, 0x3A, 0x5F), False),
    "learning_report.md": ("learning_report.docx", "Calibri", RGBColor(0x0B, 0x6E, 0x6E), False),
}


def make_reference(path, body_font, heading_colour):
    """Pandoc's default template with our fonts and colours."""
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    path.write_bytes(data)
    doc = Document(path)
    styles = {s.name.lower(): s for s in doc.styles if s.name}
    for name in ["normal", "body text", "first paragraph", "compact", "block text"]:
        if name in styles:
            styles[name].font.name = body_font
            styles[name].font.size = Pt(11)
    for level in range(1, 7):
        st = styles.get(f"heading {level}")
        if st is not None:
            st.font.name = "Calibri"
            st.font.color.rgb = heading_colour
            st.font.size = Pt({1: 18, 2: 15, 3: 13}.get(level, 12))
            st.font.bold = True
    if "title" in styles:
        styles["title"].font.name = "Calibri"
        styles["title"].font.color.rgb = heading_colour
    doc.save(path)


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_table_borders(table, colour="A6A6A6"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tbl_pr.append(borders)


def style_tables(path, header_fill):
    doc = Document(path)
    for table in doc.tables:
        set_table_borders(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell in table.rows[0].cells:
            set_cell_shading(cell, header_fill)
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9.5)
    doc.save(path)


def highlight_review_notes(path):
    """Highlight every '[REVIEW]' paragraph in yellow so it stands out in Word."""
    doc = Document(path)
    count = 0
    for par in doc.paragraphs:
        if par.text.lstrip().startswith("[REVIEW]"):
            for run in par.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            count += 1
    doc.save(path)
    return count


def main():
    for md, (out, font, colour, toc) in DOCS.items():
        ref = REPORTS / f"_reference_{out}"
        make_reference(ref, font, colour)
        cmd = ["pandoc", str(REPORTS / md), "-o", str(REPORTS / out),
               "--from", "markdown+tex_math_dollars+pipe_tables+raw_html",
               "--resource-path", str(REPORTS),
               "--reference-doc", str(ref)]
        if toc:
            cmd += ["--toc", "--toc-depth=1"]
        subprocess.run(cmd, check=True)
        style_tables(REPORTS / out, "DCE6F0" if "journal" in out else "D9EFEF")
        n_review = highlight_review_notes(REPORTS / out)
        ref.unlink()
        print(f"wrote {REPORTS / out} ({n_review} review notes highlighted)")


if __name__ == "__main__":
    main()
